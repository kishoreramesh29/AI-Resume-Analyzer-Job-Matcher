from pathlib import Path
from io import BytesIO
import re

def extract_text_from_pdf(file_bytes: bytes) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)

def extract_text_from_docx(file_bytes: bytes) -> str:
    from docx import Document
    doc = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)

def extract_text_from_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")

def extract_resume_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    elif name.endswith(".txt"):
        text = extract_text_from_txt(data)
    else:
        raise ValueError("Unsupported file type. Upload PDF, DOCX or TXT.")

    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()

def extract_contact_info(text: str) -> dict:
    email = re.findall(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", text)
    phone = re.findall(r"(?:\+?\d[\d\s().-]{8,}\d)", text)
    urls = re.findall(r"(?:https?://|www\.)\S+", text, flags=re.I)
    return {
        "email": email[0] if email else "",
        "phone": phone[0].strip() if phone else "",
        "links": urls[:5],
    }
